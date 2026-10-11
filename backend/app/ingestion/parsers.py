"""Format-specific parsers producing a common line/table representation.

Every parser yields `ParsedDocument.lines` in reading order with page numbers and
typographic hints (font size, bold) that the structure detector uses for headings.
Tables are kept whole (never split mid-row) and also exposed as records.
"""

from __future__ import annotations

import csv
import html
import io
import re
import unicodedata
from dataclasses import dataclass, field
from typing import Any

import pymupdf

from app.ingestion.validation import FileKind

_TAG = re.compile(r"<[^>]+>")
_BLOCK_TAGS = re.compile(r"</?(p|div|h[1-6]|li|tr|br|section|article|table)[^>]*>", re.I)
_HEADING_TAG = re.compile(r"<h([1-6])[^>]*>(.*?)</h\1>", re.I | re.S)
_SCRIPT_STYLE = re.compile(r"<(script|style)[^>]*>.*?</\1>", re.I | re.S)
# A page with fewer extracted characters than this is treated as scanned (OCR fallback).
_MIN_TEXT_CHARS_PER_PAGE = 40


@dataclass
class Line:
    text: str
    page: int
    font_size: float = 0.0
    bold: bool = False
    heading_level: int | None = None  # explicit from DOCX/HTML heading styles
    # Vertical position on the page (0 = top, 1 = bottom); None when unknown.
    y_ratio: float | None = None

    def __post_init__(self) -> None:
        # NFKC folds ligatures (e.g. U+FB01 'fi') and odd spaces so quotes match exactly.
        self.text = " ".join(unicodedata.normalize("NFKC", self.text).split())


@dataclass
class Table:
    page: int
    rows: list[list[str]]

    def to_markdown(self) -> str:
        if not self.rows:
            return ""
        header, *body = self.rows
        out = ["| " + " | ".join(header) + " |", "|" + "---|" * len(header)]
        out += ["| " + " | ".join(r) + " |" for r in body]
        return "\n".join(out)


@dataclass
class ParsedDocument:
    kind: FileKind
    lines: list[Line] = field(default_factory=list)
    tables: list[Table] = field(default_factory=list)
    # Spreadsheet/CSV sheets as header->value records.
    records: dict[str, list[dict[str, Any]]] = field(default_factory=dict)
    page_count: int = 1
    ocr_pages: list[int] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    @property
    def body_font_size(self) -> float:
        sizes = [round(line.font_size, 1) for line in self.lines if line.font_size]
        return max(set(sizes), key=sizes.count) if sizes else 0.0


def parse(kind: FileKind, data: bytes) -> ParsedDocument:
    if kind is FileKind.PDF:
        return _parse_pdf(data)
    if kind is FileKind.DOCX:
        return _parse_docx(data)
    if kind is FileKind.XLSX:
        return _parse_xlsx(data)
    if kind is FileKind.CSV:
        return _parse_csv(data)
    if kind is FileKind.HTML:
        return _parse_html(data.decode("utf-8"))
    return _parse_text(data.decode("utf-8"))


def _parse_pdf(data: bytes) -> ParsedDocument:
    parsed = ParsedDocument(kind=FileKind.PDF)
    with pymupdf.open(stream=data, filetype="pdf") as doc:
        parsed.page_count = doc.page_count
        for page_index, page in enumerate(doc, start=1):
            page_dict: dict[str, Any] = page.get_text("dict", sort=True)
            height = float(page.rect.height) or 1.0
            page_chars = 0
            for block in page_dict.get("blocks", []):
                for line in block.get("lines", []):
                    spans = [s for s in line.get("spans", []) if s.get("text", "").strip()]
                    if not spans:
                        continue
                    text = "".join(s["text"] for s in spans).strip()
                    page_chars += len(text)
                    parsed.lines.append(
                        Line(
                            text=text,
                            page=page_index,
                            font_size=max(float(s.get("size", 0.0)) for s in spans),
                            # PyMuPDF flag bit 4 (16) = bold
                            bold=all(int(s.get("flags", 0)) & 16 for s in spans),
                            y_ratio=float(line["bbox"][1]) / height,
                        )
                    )
            try:
                for table in page.find_tables().tables:
                    rows = [[(c or "").strip() for c in row] for row in table.extract()]
                    if rows and len(rows[0]) > 1:
                        parsed.tables.append(Table(page=page_index, rows=rows))
            except Exception:
                parsed.warnings.append(f"table detection failed on page {page_index}")
            if page_chars < _MIN_TEXT_CHARS_PER_PAGE:
                _ocr_page(page, page_index, parsed)
    return parsed


def _ocr_page(page: pymupdf.Page, page_index: int, parsed: ParsedDocument) -> None:
    """Prototype OCR fallback via PyMuPDF + Tesseract (available in the Docker image)."""
    try:
        textpage = page.get_textpage_ocr(language="eng", dpi=200, full=True)
        text = page.get_text("text", textpage=textpage)
    except Exception:
        parsed.warnings.append(f"page {page_index}: no text layer and OCR unavailable")
        return
    lines = [t.strip() for t in text.splitlines() if t.strip()]
    if lines:
        parsed.ocr_pages.append(page_index)
        parsed.lines.extend(Line(text=t, page=page_index) for t in lines)


def _parse_docx(data: bytes) -> ParsedDocument:
    import docx  # local import: heavy module, only needed for DOCX

    parsed = ParsedDocument(kind=FileKind.DOCX)
    document = docx.Document(io.BytesIO(data))
    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name if para.style is not None else "") or ""
        level = None
        if style.lower().startswith("heading"):
            digits = "".join(ch for ch in style if ch.isdigit())
            level = int(digits) if digits else 1
        elif style.lower() == "title":
            level = 0
        bold = bool(para.runs) and all(r.bold for r in para.runs if r.text.strip())
        parsed.lines.append(Line(text=text, page=1, bold=bold, heading_level=level))
    for table in document.tables:
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        if rows:
            parsed.tables.append(Table(page=1, rows=rows))
    return parsed


def _parse_xlsx(data: bytes) -> ParsedDocument:
    import openpyxl

    parsed = ParsedDocument(kind=FileKind.XLSX)
    workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    for sheet in workbook.worksheets:
        rows = [["" if v is None else str(v).strip() for v in row] for row in sheet.iter_rows(values_only=True)]
        rows = [r for r in rows if any(r)]
        if not rows:
            continue
        parsed.tables.append(Table(page=1, rows=rows))
        parsed.records[sheet.title] = _records(rows)
        for row in rows:
            parsed.lines.append(Line(text=" | ".join(c for c in row if c), page=1))
    workbook.close()
    return parsed


def _parse_csv(data: bytes) -> ParsedDocument:
    parsed = ParsedDocument(kind=FileKind.CSV)
    rows = [r for r in csv.reader(io.StringIO(data.decode("utf-8"))) if any(r)]
    if rows:
        parsed.tables.append(Table(page=1, rows=rows))
        parsed.records["csv"] = _records(rows)
        parsed.lines.extend(Line(text=" | ".join(r), page=1) for r in rows)
    return parsed


def _records(rows: list[list[str]]) -> list[dict[str, Any]]:
    header = [h.strip() for h in rows[0]]
    return [{header[i]: row[i] if i < len(row) else "" for i in range(len(header)) if header[i]} for row in rows[1:]]


def _parse_html(markup: str) -> ParsedDocument:
    parsed = ParsedDocument(kind=FileKind.HTML)
    markup = _SCRIPT_STYLE.sub(" ", markup)
    headings = {html.unescape(_TAG.sub("", m.group(2))).strip(): int(m.group(1)) for m in _HEADING_TAG.finditer(markup)}
    text = html.unescape(_TAG.sub("", _BLOCK_TAGS.sub("\n", markup)))
    for raw in text.splitlines():
        line = " ".join(raw.split())
        if line:
            parsed.lines.append(Line(text=line, page=1, heading_level=headings.get(line)))
    return parsed


def _parse_text(content: str) -> ParsedDocument:
    parsed = ParsedDocument(kind=FileKind.TEXT)
    for raw in content.splitlines():
        line = raw.strip()
        if not line:
            continue
        level = None
        if line.startswith("#"):
            level = len(line) - len(line.lstrip("#"))
            line = line.lstrip("#").strip()
        parsed.lines.append(Line(text=line, page=1, heading_level=level))
    return parsed
